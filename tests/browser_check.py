"""Real-browser check of the order support page: drives it in headless Edge (or Chrome) and reads the clipboard.

Needs Microsoft Edge or Google Chrome. Uses nothing beyond what requirements.txt already installs.
Runs the page against a temporary copy of the workbook on its own port (8765), so the real workbook and a
running copy of the app are not touched.

Usage: .venv\\Scripts\\python.exe tests\\browser_check.py [--screenshots FOLDER]
"""
import base64
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.request

import openpyxl
from websockets.sync.client import connect

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOOK = os.path.join(ROOT, "data", "stark_orders_demo.xlsx")
PORT, DEBUG_PORT = 8765, 9333
# Runs app.py with the workbook path pointed at the temporary copy.
WRAPPER = """import os, runpy, sys
sys.path.insert(0, os.getcwd())
from stark_support import ui_adapter as ui
ui.WORKBOOK_PATH = os.environ["STARK_TMP_BOOK"]
runpy.run_path(os.path.join(os.getcwd(), "app.py"), run_name="__main__")
"""


def find_browser():
    for base in (os.environ.get("ProgramFiles(x86)", ""), os.environ.get("ProgramFiles", ""), os.environ.get("LocalAppData", "")):
        for tail in (r"Microsoft\Edge\Application\msedge.exe", r"Google\Chrome\Application\chrome.exe"):
            exe = os.path.join(base, tail)
            if base and os.path.exists(exe):
                return exe
    return shutil.which("msedge") or shutil.which("chrome")


def saved_orders(path):
    """(new order number, original order number) for every row saved in New Orders."""
    ws = openpyxl.load_workbook(path)["New Orders"]
    return [(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(2, 200) if ws.cell(r, 1).value not in (None, "")]


def sha(path):
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


class Browser:
    """The few DevTools-protocol calls this check needs."""

    def __init__(self, ws):
        self.ws, self.n = ws, 0

    def cdp(self, method, **params):
        self.n += 1
        self.ws.send(json.dumps({"id": self.n, "method": method, "params": params}))
        while True:
            msg = json.loads(self.ws.recv())
            if msg.get("id") == self.n:
                if "error" in msg:
                    raise RuntimeError(msg["error"])
                return msg["result"]

    def js(self, expr):
        r = self.cdp("Runtime.evaluate", expression=expr, awaitPromise=True, returnByValue=True, userGesture=True)
        if "exceptionDetails" in r:
            raise RuntimeError(r["exceptionDetails"].get("exception", {}).get("description", "script error"))
        return r["result"].get("value")

    def wait(self, expr, seconds=30):
        end = time.time() + seconds
        while time.time() < end:
            try:
                if self.js(expr):
                    return True
            except RuntimeError:
                pass
            time.sleep(0.3)
        return False

    def text(self):
        return self.js("document.body.innerText")

    def key(self, name, code):
        for kind in ("keyDown", "keyUp"):
            self.cdp("Input.dispatchKeyEvent", type=kind, key=name, code=name, windowsVirtualKeyCode=code)

    def click(self, selector):
        self.wait(f"!!document.querySelector({selector!r})", 15)
        time.sleep(0.3)
        box = self.js(f"(() => {{ const e = document.querySelector({selector!r}); if (!e) return null; "
                      f"e.scrollIntoView({{block: 'center'}}); const r = e.getBoundingClientRect(); "
                      f"return [r.x + r.width / 2, r.y + r.height / 2]; }})()")
        if not box:
            raise RuntimeError(f"not on the page: {selector}")
        for kind in ("mousePressed", "mouseReleased"):
            self.cdp("Input.dispatchMouseEvent", type=kind, x=box[0], y=box[1], button="left", clickCount=1)
        time.sleep(0.4)

    def type_text(self, key, text):
        self.click(f".st-key-{key} input")
        self.js(f"document.querySelector('.st-key-{key} input').select()")
        self.cdp("Input.insertText", text=text)
        self.key("Enter", 13)
        time.sleep(0.6)

    def type_date(self, key, digits):
        """digits = YYYYMMDD, typed into the year / month / day parts of the date field."""
        self.js(f"document.querySelector('.st-key-{key} [role=spinbutton]').focus()")
        for ch in digits:
            self.cdp("Input.dispatchKeyEvent", type="keyDown", key=ch, code=f"Digit{ch}", text=ch,
                     windowsVirtualKeyCode=48 + int(ch))
            self.cdp("Input.dispatchKeyEvent", type="keyUp", key=ch, code=f"Digit{ch}", windowsVirtualKeyCode=48 + int(ch))
            time.sleep(0.05)
        self.key("Enter", 13)
        time.sleep(0.8)
        self.key("Escape", 27)      # close the calendar
        self.click("h1")

    def choose_inquiry(self, words):
        self.click(".st-key-inquiry input")
        self.wait("document.querySelectorAll('[role=option]').length > 0", 5)
        self.js(f"[...document.querySelectorAll('[role=option]')].find(o => o.innerText.includes({words!r})).click()")
        time.sleep(1)

    def value(self, key, tag="input"):
        return self.js(f"(document.querySelector('.st-key-{key} {tag}') || {{}}).value")

    def resize(self, width, height=1300):
        self.cdp("Emulation.setDeviceMetricsOverride", width=width, height=height, deviceScaleFactor=1, mobile=False)
        time.sleep(1)

    def button_state(self, key):
        """How a button really looks: size, on screen, not covered, label on one line and not cut off, readable colours."""
        return self.js(f"""(() => {{
            const e = document.querySelector('.st-key-{key} button');
            if (!e) return null;
            e.scrollIntoView({{block: 'center'}});
            const r = e.getBoundingClientRect(), c = getComputedStyle(e);
            const top = document.elementFromPoint(r.x + r.width / 2, r.y + r.height / 2);
            return {{label: e.innerText.trim(), left: r.left, right: r.right, top: r.top, width: r.width, height: r.height,
                     onScreen: r.top >= 0 && r.left >= 0 && r.bottom <= innerHeight && r.right <= innerWidth,
                     uncovered: !!top && (top === e || e.contains(top)),
                     cutOff: e.scrollWidth > e.clientWidth + 1 || e.scrollHeight > e.clientHeight + 1,
                     shown: c.display !== 'none' && c.visibility !== 'hidden' && parseFloat(c.opacity) > 0.9,
                     text: c.color, background: c.backgroundColor, border: c.borderTopColor}};
        }})()""")

    def clipboard(self):
        return self.js("navigator.clipboard.readText()").replace("\r", "")

    def screenshot(self, path):
        height = self.js("Math.ceil(document.querySelector('[data-testid=stMain]')?.scrollHeight || 1300)")
        self.cdp("Emulation.setDeviceMetricsOverride", width=1500, height=min(max(height + 40, 900), 4000),
                 deviceScaleFactor=1, mobile=False)
        time.sleep(0.8)
        with open(path, "wb") as fh:
            fh.write(base64.b64decode(self.cdp("Page.captureScreenshot", format="png")["data"]))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    shots = sys.argv[sys.argv.index("--screenshots") + 1] if "--screenshots" in sys.argv else None
    theme = "dark" if "--dark" in sys.argv else "light"
    exe = find_browser()
    if not exe:
        print("SKIPPED: Microsoft Edge or Google Chrome was not found on this computer.")
        return 0
    if shots:
        os.makedirs(shots, exist_ok=True)
    real_before = sha(BOOK)
    tmp = tempfile.mkdtemp(prefix="stark-browser-")
    book = os.path.join(tmp, "book.xlsx")
    wb = openpyxl.load_workbook(BOOK)
    for row in wb["New Orders"].iter_rows(min_row=2):   # start with no saved new orders, whatever the live workbook holds
        for cell in row:
            cell.value = None
    wb["New Orders"]["N1"] = wb["New Orders"]["O1"] = None   # ...and no recorded "highest order number ever issued"
    wb.save(book)
    wrapper = os.path.join(tmp, "page.py")
    with open(wrapper, "w", encoding="utf-8") as fh:
        fh.write(WRAPPER)
    before = sha(book)
    quiet = {"stdout": subprocess.DEVNULL, "stderr": subprocess.DEVNULL}
    server = subprocess.Popen(
        [sys.executable, "-m", "streamlit", "run", wrapper, "--server.address", "localhost", "--server.port", str(PORT),
         "--server.headless", "true", "--browser.gatherUsageStats", "false"],
        cwd=ROOT, env=dict(os.environ, STARK_TMP_BOOK=book, PYTHONUTF8="1"), **quiet)
    browser = subprocess.Popen(
        [exe, "--headless=new", f"--remote-debugging-port={DEBUG_PORT}", f"--user-data-dir={os.path.join(tmp, 'profile')}",
         "--window-size=1500,1300", "--no-first-run", "--remote-allow-origins=*", "about:blank"], **quiet)
    results = []

    def note(name, ok, detail=""):
        results.append(bool(ok))
        print(("PASS  " if ok else "FAIL  ") + name + (f" | {detail}" if detail and not ok else ""))

    try:
        tabs = None
        for _ in range(80):
            try:
                urllib.request.urlopen(f"http://localhost:{PORT}/_stcore/health", timeout=1)
                tabs = json.load(urllib.request.urlopen(f"http://localhost:{DEBUG_PORT}/json", timeout=1))
                break
            except Exception:
                time.sleep(0.5)
        if not tabs:
            print("FAIL  the app or the browser did not start")
            return 1
        with connect(next(t for t in tabs if t["type"] == "page")["webSocketDebuggerUrl"], max_size=None) as ws:
            b = Browser(ws)
            b.cdp("Page.enable")
            b.cdp("Runtime.enable")
            b.cdp("Browser.grantPermissions", permissions=["clipboardReadWrite", "clipboardSanitizedWrite"],
                  origin=f"http://localhost:{PORT}")
            b.cdp("Emulation.setEmulatedMedia", features=[{"name": "prefers-color-scheme", "value": theme}])
            b.cdp("Page.navigate", url=f"http://localhost:{PORT}")

            def snap(name):
                if shots:
                    b.screenshot(os.path.join(shots, name))

            # start page
            note("page loads", b.wait("document.body.innerText.includes('Check Order and Prepare Reply')"))
            time.sleep(1)
            snap("1_start.png")
            text = b.text()
            note("no iframe, script text or traceback on the page",
                 b.js("document.querySelectorAll('iframe').length") == 0 and "Traceback" not in text and "<script" not in text)
            note("Streamlit's Deploy button is hidden", "Deploy" not in text, text[:80])

            # validation
            b.click(".st-key-check button")
            note("a missing contact date gives a plain message",
                 b.wait("document.body.innerText.includes('Enter the date the customer contacted you')"))

            # where is my order: late, not eligible
            b.type_text("ident", "1650")
            b.type_date("contact", "20261008")
            b.type_date("app", "20261008")
            b.click(".st-key-check button")
            note("status, recommended action and email appear",
                 b.wait("document.body.innerText.includes('2. Recommended action') && "
                        "!!document.querySelector('.st-key-email_body textarea')"))
            time.sleep(1)
            snap("2_late.png")
            text = b.text()
            note("late order: plain status and 'do not reship yet'",
                 "Running late, waiting for reshipment eligibility" in text and "Do not reship yet, send the reply" in text)
            note("shipping, expected delivery and contact-us dates are shown in the result",
                 "Shipped: Wednesday, September 30, 2026" in text and "Expected delivery: Tuesday, October 6, 2026" in text
                 and "Contact-us date: Tuesday, October 13, 2026" in text
                 and "If the package has not arrived by Tuesday, October 13, 2026, the customer should contact us again" in text)
            note("the email gives the same contact-us date", "Tuesday, October 13" in b.value("email_body", "textarea"))
            b.js("[...document.querySelectorAll('summary')].find(e => e.innerText.includes('Internal explanation')).click()")
            time.sleep(1)
            inner = b.text()
            note("internal explanation shows X, 2X and the extra day",
                 "X for Regular shipping to Johannesburg is 4 business days" in inner
                 and "2X is 8 business days after the shipping date: Monday, October 12, 2026." in inner
                 and "day 8 Mon Oct 12, 2026" in inner
                 and "Plus 1 calendar day gives the contact-us date: Tuesday, October 13, 2026." in inner)
            snap("2b_explanation.png")
            note("the email itself has no formula or shipping-method wording",
                 not any(w in b.value("email_body", "textarea") for w in ("2X", "business days", "Regular", "Express")))
            subject, body = b.value("email_subject"), b.value("email_body", "textarea")
            b.click(".st-key-copy_email button")
            time.sleep(1)
            copied = b.clipboard()
            note("Copy Email copies the subject and the whole email",
                 copied == f"Subject: {subject}\n\n{body}" and copied.startswith("Subject: Your order 1650: delivery update\n\nHi M'Baku,")
                 and copied.rstrip().endswith("Stark Industries Support"), repr(copied[:80]))
            note("a 'copied' confirmation is shown", "Email copied" in b.text())
            b.click(".st-key-email_body textarea")
            b.js("(() => { const t = document.querySelector('.st-key-email_body textarea'); "
                 "t.setSelectionRange(t.value.length, t.value.length); })()")
            b.cdp("Input.insertText", text="\nSam")
            b.click(".st-key-copy_email button")
            time.sleep(1.5)
            note("an edit made just before clicking is included in the copy",
                 b.clipboard() == f"Subject: {subject}\n\n{body}\nSam")

            # refund
            b.choose_inquiry("refund")
            note("changing the inquiry flags the old result", "You changed the inquiry" in b.text())
            b.type_text("ident", "1814")
            b.click(".st-key-check button")
            note("refund inquiry gives the refund answer",
                 b.wait("(document.querySelector('.st-key-email_body textarea') || {}).value?.includes('for a refund')"))
            time.sleep(1)
            snap("3_refund.png")
            text = b.text()
            note("refund: amount shown as plain text, voucher only after the return",
                 "Refund allowed: $45 voucher after the return" in text and "only after the customer replies" in text)
            note("refund preview wrote nothing", sha(book) == before)

            # exchange: allowed, but no order until the tracking number arrives
            b.choose_inquiry("exchange an item")
            b.click(".st-key-new_item input")
            b.wait("document.querySelectorAll('[role=option]').length > 0", 5)
            b.js("[...document.querySelectorAll('[role=option]')].find(o => o.innerText.includes('Vibranium Plate (Red)')).click()")
            time.sleep(1)
            b.click(".st-key-check button")
            note("exchange inquiry gives the exchange answer",
                 b.wait("(document.querySelector('.st-key-email_body textarea') || {}).value?.includes('Yes, we can exchange')"))
            time.sleep(1)
            snap("4_exchange.png")
            text = b.text()
            note("exchange: allowed, and no order is offered for creation yet",
                 "Exchange allowed" in text and "Do not create the exchange order yet" in text
                 and not b.js("!!document.querySelector('.st-key-create button')"))
            note("exchange preview wrote nothing", sha(book) == before)

            # late / not arrived: eligible, preview, confirm, save
            b.choose_inquiry("late or")
            b.type_text("ident", "1650")
            b.type_date("contact", "20261030")
            b.click(".st-key-check button")
            note("eligible order: preview shown and the email held back",
                 b.wait("document.body.innerText.includes('preview only, not saved') && "
                        "document.body.innerText.includes('The email appears after you create the order')")
                 and not b.js("!!document.querySelector('.st-key-email_body textarea')"))
            time.sleep(1)
            snap("5_preview.png")
            note("replacement is by Regular shipping", "Create the replacement order (Regular shipping)" in b.text())
            note("preview wrote nothing", sha(book) == before)
            text = b.text()
            note("preview lists the new order number, item, ship, delivery and contact-us dates",
                 all(line in text for line in ("New order number: 5021", "Item: 1 × vibranium plate, red",
                                               "Ships: Friday, October 30, 2026", "Expected delivery: Thursday, November 5, 2026",
                                               "Contact-us date: Thursday, November 12, 2026",
                                               "Shipping method (internal, not in the email): Regular")), text[-900:])
            b.click(".st-key-create button")
            note("a confirm step appears", b.wait("document.body.innerText.includes('Save this replacement order to the workbook?')"))
            note("the confirm step wrote nothing", sha(book) == before)
            def buttons_visible(yes_key, yes_label, no_key, picture):
                for width in (1500, 1100, 800, 420):
                    b.resize(width)
                    ok_btn, no_btn = b.button_state(yes_key), b.button_state(no_key)
                    good = bool(ok_btn and no_btn) and all(
                        x["shown"] and x["onScreen"] and x["uncovered"] and not x["cutOff"] and x["height"] <= 60
                        and x["text"] != x["background"] for x in (ok_btn, no_btn))
                    good = good and ok_btn["label"] == yes_label and no_btn["label"] == "Cancel" \
                        and ok_btn["width"] >= 100 and no_btn["width"] >= 60 \
                        and (ok_btn["right"] <= no_btn["left"] or ok_btn["top"] != no_btn["top"])   # side by side, not overlapping
                    note(f"{yes_label} / Cancel are fully visible at {width}px wide ({theme} theme)", good,
                         json.dumps({"yes": ok_btn, "cancel": no_btn}))
                    if shots:
                        with open(os.path.join(shots, f"{picture}_{width}.png"), "wb") as fh:
                            fh.write(base64.b64decode(b.cdp("Page.captureScreenshot", format="png")["data"]))
                b.resize(1500)

            buttons_visible("save", "Confirm and save", "cancel", "5b_confirm")
            b.click(".st-key-cancel button")
            note("Cancel: says nothing was saved and offers Create order again",
                 b.wait("document.body.innerText.includes('Cancelled. Nothing was saved.')")
                 and b.js("!!document.querySelector('.st-key-create button') && !document.querySelector('.st-key-save button')"))
            note("Cancel wrote nothing", sha(book) == before)
            b.click(".st-key-create button")
            b.wait("!!document.querySelector('.st-key-save button')")
            b.click(".st-key-save button")
            note("after Confirm and save: order number and email appear",
                 b.wait("document.body.innerText.includes('Done: order 5021 created') && (document.querySelector("
                        "'.st-key-email_body textarea') || {}).value?.includes('Your replacement order number is 5021')"))
            mail = b.value("email_body", "textarea") or ""
            note("the replacement email gives ship, delivery and contact-us dates but not the shipping method",
                 "scheduled to ship on Friday, October 30 and should arrive by Thursday, November 5" in mail
                 and "If it hasn't arrived by Thursday, November 12, please contact us again" in mail
                 and "Regular" not in mail and "regular shipping" not in mail, mail)
            time.sleep(1)
            snap("6_saved.png")
            saved = sha(book)
            note("the workbook copy changed", saved != before)
            note("status now says the replacement was sent", "Replacement already sent (order 5021)" in b.text())
            note("no save buttons remain after saving",
                 not b.js("!!document.querySelector('.st-key-save button') || !!document.querySelector('.st-key-create button')"))
            b.click(".st-key-check button")
            note("repeating the inquiry creates nothing and points to order 5021",
                 b.wait("document.body.innerText.includes('order 5021 already covers this')") and sha(book) == saved
                 and not b.js("!!document.querySelector('.st-key-create button')"))

            # demonstrations: clear the saved replacement, then the original can be processed again
            note("the clear option sits in its own demonstration section, closed by default",
                 "Demonstration and testing: clear a saved replacement" in b.text()
                 and b.js("[...document.querySelectorAll('details')].find(d => d.innerText.includes("
                          "'Demonstration and testing')).open") is False)
            b.js("[...document.querySelectorAll('summary')].find(e => e.innerText.includes('Demonstration and testing')).click()")
            time.sleep(1)
            b.click(".st-key-clear_5021 button")
            asked = b.wait("document.body.innerText.includes('Clear saved replacement order 5021?')")
            text = b.text()
            note("clearing asks first and names both orders",
                 asked and "Original order: 1650" in text and "Saved replacement order: 5021" in text
                 and "can be processed again" in text and saved_orders(book) == [(5021, 1650)])
            snap("6b_clear_confirm.png")
            buttons_visible("confirm_clear", "Confirm clear", "cancel_clear", "6c_clear")
            b.click(".st-key-cancel_clear button")
            note("Cancel clear: nothing changed",
                 b.wait("document.body.innerText.includes('Cancelled. Nothing was changed.')") and sha(book) == saved
                 and saved_orders(book) == [(5021, 1650)])
            b.click(".st-key-clear_5021 button")
            b.wait("!!document.querySelector('.st-key-confirm_clear button')")
            b.click(".st-key-confirm_clear button")
            note("Confirm clear: success message, and the order is eligible for reshipment again",
                 b.wait("document.body.innerText.includes('Saved replacement order 5021 was cleared from the workbook')")
                 and "Eligible for reshipment" in b.text() and "Replacement already sent" not in b.text())
            time.sleep(1)
            snap("6d_cleared.png")
            note("the replacement row is gone and no new order was saved", saved_orders(book) == [])
            note("a new replacement is only previewed: Create order is offered, no email yet",
                 "preview only, not saved" in b.text() and b.js("!!document.querySelector('.st-key-create button')")
                 and not b.js("!!document.querySelector('.st-key-email_body textarea')"))
            note("the preview shows a new number, 5022, not the cleared 5021",
                 "New order number: 5022" in b.text() and "New order number: 5021" not in b.text())
            b.click(".st-key-create button")
            b.wait("!!document.querySelector('.st-key-save button')")
            note("still nothing saved before Confirm and save", saved_orders(book) == [])
            b.click(".st-key-save button")
            note("the normal workflow saves a new replacement under the new number 5022",
                 b.wait("document.body.innerText.includes('Replacement order saved.')") and saved_orders(book) == [(5022, 1650)]
                 and "Your replacement order number is 5022." in (b.value("email_body", "textarea") or ""))
            saved = sha(book)

            # check another order
            b.click(".st-key-another button")
            cleared = b.wait("document.body.innerText.includes('Fill in the customer inquiry above')")
            time.sleep(1.5)
            snap("7_cleared.png")
            text = b.text()
            note("Check Another Order clears the fields and the result",
                 cleared and b.value("ident") == "" and b.value("inquiry") == "Where is my order?" and "yyyy" in text.lower()
                 and "Recommended action" not in text and not b.js("!!document.querySelector('.st-key-email_body textarea')"),
                 text[:300])
            note("Check Another Order left the saved order alone", sha(book) == saved)
    finally:
        for proc in (browser, server):
            proc.terminate()
        for proc in (browser, server):
            try:
                proc.wait(timeout=10)
            except subprocess.TimeoutExpired:
                proc.kill()
        for _ in range(10):     # the browser lets go of its profile folder a moment after it exits
            shutil.rmtree(tmp, ignore_errors=True)
            if not os.path.exists(tmp):
                break
            time.sleep(1)
    note("the real workbook is unchanged", sha(BOOK) == real_before)
    print(f"\n{sum(results)}/{len(results)} browser checks passed")
    return 0 if all(results) else 1


if __name__ == "__main__":
    sys.exit(main())
