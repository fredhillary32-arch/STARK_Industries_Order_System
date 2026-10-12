"""Tests for the browser interface and the customer emails: the adapter (no browser needed) and headless page runs.

Runs on temporary copies of the sanitized demo workbook; the source workbook is never modified.
Usage: .venv\\Scripts\\python.exe tests\\test_ui_adapter.py
"""
import atexit
import copy
import hashlib
import os
import shutil
import sys
import tempfile
import unittest
from datetime import date, timedelta
from unittest import mock

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from stark_support import engine as E  # noqa: E402
from stark_support import reply  # noqa: E402
from stark_support import ui_adapter as ui  # noqa: E402

BOOK = os.path.join(ROOT, "data", "stark_orders_demo.xlsx")
D = lambda m, d: date(2026, m, d)  # noqa: E731
APP_A, APP_B = D(10, 8), D(10, 30)
ORDERS = [1024, 1545, 1614, 1628, 1650, 1655, 1814, 1855, 2314, 3514, 3520, 5020]
TMP = tempfile.mkdtemp(prefix="stark-ui-tests-")
# Not tearDownModule alone: the page tests make Streamlit swap out __main__, and unittest then can't find it.
atexit.register(shutil.rmtree, TMP, ignore_errors=True)
START = {}


def sha(path):
    with open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def temp_book():
    """A temporary copy of the working workbook with no saved new orders, so tests don't depend on real transactions."""
    base = os.path.join(TMP, "base.xlsx")
    if not os.path.exists(base):
        wb = openpyxl.load_workbook(BOOK)
        demo_emails = {
            2: "multi-order-c@example.com", 3: "multi-order-a@example.com",
            5: "multi-order-a@example.com", 6: "single-order@example.com",
            7: "multi-order-b@example.com", 9: "multi-order-b@example.com",
            10: "multi-order-b@example.com", 11: "multi-order-c@example.com",
            12: "multi-order-a@example.com",
        }
        for row, email in demo_emails.items():
            wb["Orders"].cell(row, 2).value = email
        demo_names = {
            2: "Norman Osbourne", 3: "Clint Barton", 5: "Clint Barton",
            6: "M'Baku", 7: "Scott Lang", 8: "Nat", 9: "Scott Lang", 10: "Scott Lang",
            11: "Norman Osbourne", 12: "Bruce Banner", 13: "Miles Morales",
        }
        for row, name in demo_names.items():
            wb["Orders"].cell(row, 3).value = name
        for row in wb["New Orders"].iter_rows(min_row=2):
            for cell in row:
                cell.value = None
        wb["New Orders"]["N1"] = wb["New Orders"]["O1"] = None    # and no recorded "highest order number ever issued"
        wb.save(base)
    p = os.path.join(TMP, f"book{len(os.listdir(TMP))}.xlsx")
    shutil.copy(base, p)
    return p


def new_order_rows(path):
    ws = openpyxl.load_workbook(path)["New Orders"]
    return [r for r in range(2, 200) if ws[f"A{r}"].value not in (None, "")]


def save(result, path):
    """Create the previewed order the way the page does."""
    state = {}
    return ui.save_once(state, ui.remember_pending(state, result), path=path)


def setUpModule():
    START["book"] = sha(BOOK)


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


class CheckOrder(unittest.TestCase):
    def test_1650_preview_reship_writes_nothing(self):
        p = temp_book()
        before = sha(p)
        r = ui.check_order("1650", D(10, 30), APP_A, path=p)
        self.assertTrue(r["ok"])
        self.assertEqual(r["analysis"]["state"], E.RESHIP)
        self.assertEqual(len(r["previews"]), 1)
        row = r["previews"][0]
        self.assertEqual((row["method"], row["ship"], row["est"], row["contact"]),
                         ("Regular", D(10, 30), D(11, 5), D(11, 12)))       # due Thu Nov 5, contact-us Thu Nov 12
        self.assertEqual(r["saved_orders"], [])
        self.assertTrue(r["email_withheld"])
        self.assertEqual(r["email"], "")
        self.assertEqual(r["linked_orders"], [])                           # the preview is not listed as a real order
        self.assertEqual(sha(p), before)
        self.assertEqual(new_order_rows(p), [])

    def test_1650_on_app_date_oct_30_is_late_not_eligible(self):
        p = temp_book()
        before = sha(p)
        r = ui.check_order("1650", D(10, 30), APP_B, path=p)
        dates = dict(r["key_dates"])
        self.assertEqual(r["analysis"]["state"], E.LATE)
        self.assertEqual(dates["Expected delivery"], "Wed Oct 28, 2026")
        self.assertEqual(dates["Contact-us date"], "Wed Nov 4, 2026")
        self.assertEqual(dates["Order date (rolling)"], "Wed Oct 21, 2026")
        self.assertEqual(r["previews"], [])
        self.assertIn("Wednesday, November 4", r["email"])
        self.assertEqual(sha(p), before)

    def test_not_due_shows_not_late(self):
        r = ui.check_order("5020", D(10, 12), APP_A, path=temp_book())
        self.assertEqual(r["analysis"]["state"], E.NOT_YET_DUE)
        self.assertEqual(dict(r["key_dates"])["Days late"], "Not late")

    def test_timeline_places_the_contact_date(self):
        r = ui.check_order("1650", D(10, 8), APP_A, path=temp_book())
        self.assertEqual([(e["date"], e["label"]) for e in r["timeline"]], [
            (D(9, 29), "Ordered (rolling)"), (D(9, 30), "Shipped (rolling)"), (D(10, 6), "Expected delivery"),
            (D(10, 8), "Customer contacted us"), (D(10, 13), "Contact-us date")])

    def test_email_shared_by_several_orders_returns_matches(self):
        p = temp_book()
        for email, numbers in (("multi-order-a@example.com", [1545, 1628, 3520]),
                               ("multi-order-b@example.com", [1655, 1855, 2314]),
                               ("multi-order-c@example.com", [1024, 3514])):
            r = ui.check_order(email, D(10, 8), APP_A, path=p)
            self.assertTrue(r["ok"])
            self.assertEqual(r["decision"], "AMBIGUOUS")
            self.assertEqual(r["matches"], numbers)
            self.assertEqual([m["number"] for m in r["match_details"]], numbers)
            self.assertIsNone(r["order"])

    def test_email_with_one_order_is_found(self):
        r = ui.check_order("single-order@example.com", D(10, 8), APP_A, path=temp_book())
        self.assertEqual(r["order"]["number"], 1650)

    def test_unknown_order_number_is_a_clear_error(self):
        r = ui.check_order("9999", D(10, 8), APP_A, path=temp_book())
        self.assertFalse(r["ok"])
        self.assertIn("No order with the number 9999", r["error"])

    def test_missing_contact_date_is_a_clear_error(self):
        r = ui.check_order("1650", None, APP_A, path=temp_book())
        self.assertFalse(r["ok"])
        self.assertIn("date the customer contacted you", r["error"])

    def test_bad_input_and_empty_input(self):
        self.assertIn("order number", ui.check_order("M'Baku", D(10, 8), APP_A, path=temp_book())["error"])
        self.assertIn("order number", ui.check_order("", D(10, 8), APP_A, path=temp_book())["error"])
        self.assertIn("Choose what the customer is asking",
                      ui.prepare_reply(None, "1650", D(10, 8), APP_A, path=temp_book())["error"])

    def test_contact_before_order_names_both_dates(self):
        r = ui.check_order("1650", D(10, 8), APP_B, path=temp_book())
        self.assertEqual(r["decision"], "CONTACT_BEFORE_ORDER")
        self.assertIn("Thu Oct 8, 2026", r["warning"])
        self.assertIn("Fri Oct 30, 2026", r["warning"])
        self.assertEqual(r["email"], "")

    def test_unreadable_workbook_is_a_plain_error(self):
        with self.assertLogs("stark_support.ui", level="ERROR"):     # the traceback goes to the console log
            r = ui.check_order("1650", D(10, 8), APP_A, path=os.path.join(TMP, "missing.xlsx"))
        self.assertFalse(r["ok"])
        self.assertIn("Something went wrong", r["error"])


class Actions(unittest.TestCase):
    def test_refund_preview_uses_engine_amounts_and_writes_nothing(self):
        p = temp_book()
        before = sha(p)
        r = ui.run_action("refund", 1814, D(10, 8), APP_A, path=p, item=1, qty=1, ask_remaining=True)
        self.assertEqual(r["decision"], "REFUND_OK")
        self.assertEqual(r["internal"]["facts"]["refund_amount"], 15)
        self.assertIn("$15 voucher", r["email"])
        self.assertEqual(sha(p), before)

    def test_damaged_preview_withholds_the_email(self):
        p = temp_book()
        before = sha(p)
        r = ui.run_action("damaged", 3514, D(10, 8), APP_A, path=p)
        self.assertEqual(r["previews"][0]["method"], "Express")
        self.assertTrue(r["email_withheld"])
        self.assertEqual(sha(p), before)

    def test_force_new_is_not_available(self):
        for call in (ui.run_action("lost", 1650, D(10, 30), APP_A, path=temp_book(), force_new=True),
                     ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=temp_book(), force_new=True)):
            self.assertFalse(call["ok"])
            self.assertIn("force_new", call["error"])

    def test_lint_flags_internal_terms(self):
        self.assertEqual(ui.lint_email("Your order is on its way."), [])
        self.assertEqual(ui.lint_email("Your order is lost after 2X."), ["2X", "lost"])


class Emails(unittest.TestCase):
    """The email answers the customer's actual question, using the dates and decisions the engine produced."""

    def ask(self, kind, order, on, app=APP_A, path=None, **fields):
        return ui.prepare_reply(kind, str(order), on, app, path=path or temp_book(), **fields)

    def test_different_questions_get_different_explicit_answers(self):
        p = temp_book()
        where = self.ask("status", 5020, D(10, 12), path=p)["email"]
        late = self.ask("lost", 5020, D(10, 12), path=p)["email"]
        cancel = self.ask("cancel", 5020, D(10, 12), path=p)["email"]
        refund = self.ask("refund", 5020, D(10, 12), path=p)["email"]
        exchange = self.ask("exchange", 5020, D(10, 12), path=p, new_item="Waterproof Spray (Medium)")["email"]
        self.assertEqual(len({where, late, cancel, refund, exchange}), 5)
        self.assertIn("Your order 5020 is on its way.", where)
        self.assertIn("so it isn't late yet", late)
        self.assertIn("we aren't able to cancel or change an order once it's been placed", cancel)
        self.assertIn("Yes, you can return the Waterproof Spray (Large) for a refund.", refund)
        self.assertIn("Yes, we can exchange the Waterproof Spray (Large) for the Waterproof Spray (Medium).", exchange)
        for text in (where, late, cancel, refund, exchange):       # the answer is in the first two paragraphs
            self.assertTrue(text.startswith("Hi Miles,\n\n"))
            self.assertTrue(text.endswith("\n\nBest,\nStark Industries Support"))

    def test_wording_matches_each_order_status(self):
        want = {D(10, 5): (E.NOT_YET_DUE, "is on its way", "On schedule"),
                D(10, 6): (E.DUE, "is due to arrive today", "On schedule, due today"),
                D(10, 8): (E.LATE, "is running late", "Running late, waiting for reshipment eligibility")}
        for on, (state, phrase, label) in want.items():
            r = self.ask("status", 1650, on)
            self.assertEqual(r["decision"], state)
            self.assertIn(phrase, r["email"])
            self.assertEqual(r["status"]["label"], label)
        r = self.ask("status", 5020, D(10, 8))
        self.assertIn("hasn't shipped yet", r["email"])
        self.assertEqual(r["status"]["label"], "Not shipped yet")

    def test_dates_in_the_email_are_the_calculated_dates(self):
        r = self.ask("status", 1650, D(10, 8))
        f = r["internal"]["facts"]
        self.assertEqual((f["ship_date"], f["expected_delivery"], f["contact_us_date"]),
                         ("2026-09-30", "2026-10-06", "2026-10-13"))
        for text in ("Wednesday, September 30", "Tuesday, October 6", "Tuesday, October 13"):
            self.assertIn(text, r["email"])
        r = self.ask("status", 1650, D(10, 30), app=APP_B)          # the same order after its dates have rolled
        for text in ("Thursday, October 22", "Wednesday, October 28", "Wednesday, November 4"):
            self.assertIn(text, r["email"])

    def test_every_generated_email_passes_the_quality_check(self):
        """All 12 orders x every inquiry x several contact dates: no email states a date, amount or order number
        the engine didn't calculate, uses an internal term, or drops its greeting or closing."""
        p = temp_book()
        before = sha(p)
        base = ui.load(p, app_date=APP_A)       # read once; every case gets its own copy, as a fresh load would
        extras = {"exchange": {"new_item": "Waterproof Spray (Small)"}, "refund": {"ask_remaining": True},
                  "exchange_tracking": {"new_item": "Waterproof Spray (Small)", "tracking_number": "1ZTEST"}}
        checked = 0
        with mock.patch.object(ui, "load", side_effect=lambda path, app_date: copy.deepcopy(base)):
            for kind, _ in ui.INQUIRIES:
                for number in ORDERS:
                    for days in (-2, 0, 1, 3, 6, 12, 22, 40):
                        r = self.ask(kind, number, APP_A + timedelta(days), path=p, **extras.get(kind, {}))
                        if not r["ok"] or not r["email"]:
                            continue
                        checked += 1
                        self.assertEqual(r["email_issues"], [], f"{kind} {number} +{days}: {r['email']}")
                        self.assertNotRegex(r["email"].lower(), r"\blost\b")
                        self.assertNotRegex(r["email"], r"(?i)\b(regular|express) shipping\b")
                        self.assertTrue(r["subject"].endswith(str(number)) or f"order {number}" in r["subject"])
        self.assertGreater(checked, 300)
        self.assertEqual(sha(p), before)

    def test_ineligible_orders_get_no_promise_of_action(self):
        late = self.ask("lost", 1650, D(10, 8))                     # late, reshippable only from Oct 13
        self.assertEqual(late["previews"], [])
        self.assertIn("If it hasn't reached you by Tuesday, October 13", late["email"])     # conditional, with its date
        self.assertNotRegex(late["email"], r"new order number|I've set up|ships today")
        self.assertEqual(late["recommendation"]["title"], "Do not reship yet, send the reply")
        for r, title in ((self.ask("refund", 1655, D(10, 8)), "Refund not allowed (hoverboard)"),
                         (self.ask("refund", 1545, D(10, 5)), "Refund not allowed (too late)"),
                         (self.ask("exchange", 1650, D(10, 8), new_item="Quiver (Aluminum)"),
                          "Exchange not allowed (vibranium item)"),
                         (self.ask("damaged", 1024, D(10, 8)), "Replacement not allowed (too late)")):
            self.assertEqual(r["recommendation"]["title"], title)
            self.assertEqual(r["previews"], [])
            self.assertIn("not able to", r["email"].replace("can't be exchanged", "not able to"))
            self.assertNotRegex(r["email"], r"voucher|new order number|I've set up")

    def test_refusals_say_what_is_still_possible(self):
        late_refund = self.ask("refund", 1545, D(10, 5))            # refund ended Oct 2, exchange open until Oct 8
        self.assertIn("I can still help you exchange them", late_refund["email"])
        self.assertIn("Thursday, October 8", late_refund["email"])
        self.assertNotIn("exchange", self.ask("refund", 1545, D(10, 9))["email"])           # both windows closed
        self.assertIn("help with a return.", self.ask("cancel", 1650, D(10, 8))["email"])   # vibranium: no exchange
        self.assertIn("help with an exchange if", self.ask("cancel", 2314, D(10, 7))["email"])   # hoverboard: no return
        self.assertNotIn("happy to help", self.ask("cancel", 1024, D(10, 8))["email"])      # both windows closed

    def test_item_inquiry_before_shipping_answers_that_inquiry(self):
        for kind, phrase in (("refund", "I can't set up a refund yet because your order 5020 hasn't shipped."),
                             ("damaged", "I can't set up a replacement yet because your order 5020 hasn't shipped.")):
            r = self.ask(kind, 5020, D(10, 8))
            self.assertIn(phrase, r["email"])
            self.assertIn("Friday, October 9", r["email"])
            self.assertEqual(r["recommendation"]["title"], "Not possible yet: the order hasn't shipped")
            self.assertEqual(r["previews"], [])

    def test_no_email_until_the_order_exists_then_it_follows_the_shipping_rule(self):
        for kind, order, on, fields, method, word in (
                ("lost", 1650, D(10, 30), {}, "Regular", None),
                ("damaged", 3514, D(10, 8), {}, "Express", None),
                ("wrong", 3520, D(10, 12), {}, "Express", None),
                ("exchange_tracking", 1628, D(10, 12), {"new_item": "Waterproof Spray (Medium)", "tracking_number": "1Z9",
                                                        "offered_on": D(10, 8)}, "Regular", None)):
            p = temp_book()
            preview = self.ask(kind, order, on, path=p, **fields)
            self.assertTrue(preview["email_withheld"])
            self.assertEqual(preview["email"], "")
            self.assertEqual(preview["previews"][0]["method"], method)
            self.assertIn(f"({method} shipping)", preview["recommendation"]["title"])
            done = save(preview, p)
            row = done["saved_orders"][0]
            self.assertEqual(row["method"], method)
            self.assertIn(f"order number is {row['number']}.", done["email"])
            self.assertIn(f"scheduled to ship on {E.fmt(row['ship'])} and should arrive by {E.fmt(row['est'])}", done["email"])
            self.assertIn(f"If it hasn't arrived by {E.fmt(row['contact'])}, please contact us again", done["email"])
            self.assertNotRegex(done["email"], r"(?i)\b(regular|express)\b")     # the method stays internal
            self.assertEqual(dict(row["details"])["Shipping method (internal, not in the email)"], method)
            self.assertEqual(done["email_issues"], [])
            self.assertEqual(done["recommendation"]["title"], f"Done: order {row['number']} created")
            if kind == "lost":
                self.assertEqual(preview["status"]["label"], "Eligible for reshipment")
                self.assertEqual(done["status"]["label"], f"Replacement already sent (order {row['number']})")

    def test_quality_check_catches_bad_edits(self):
        r = self.ask("refund", 1814, D(10, 8), qty=1)
        good = r["email"]
        self.assertEqual(ui.check_email(r, good), [])
        for edited, expect in (
                (good.replace("Hi Nat,\n\n", ""), "greeting"),
                (good.replace("Best,\nStark Industries Support", ""), "closing"),
                (good.replace("$15", "$50"), "amount"),
                (good + "\nP.S. It will arrive by October 20.", "date"),
                (good.replace("refund", "thing"), "customer's question"),
                (good.replace("Yes,", "As per policy,"), "internal terms"),
                (good.replace("we'll send you a $15 voucher", "the voucher has been issued"), "already issued"),
                (good.replace("Yes,", "Your new order number is 7777. Yes,"), "order number")):
            self.assertTrue(any(expect in issue for issue in ui.check_email(r, edited)), expect)
        late = self.ask("lost", 1650, D(10, 8))
        promised = late["email"].replace("I'm sorry,", "I've set up a replacement and")
        self.assertTrue(any("no new order was saved" in i for i in ui.check_email(late, promised)))
        self.assertEqual(ui.check_email(late, late["email"].replace("Best,", "Kind regards,")), [])

    def test_subject_and_copy_text(self):
        subjects = {kind: self.ask(kind, 5020, D(10, 12), new_item="Waterproof Spray (Medium)" if "exchange" in kind else None,
                                   )["subject"] for kind in ("status", "lost", "cancel", "refund", "exchange")}
        self.assertEqual(subjects["status"], "Your order 5020: delivery update")
        self.assertEqual(subjects["refund"], "Your refund request for order 5020")
        self.assertEqual(subjects["exchange"], "Your exchange request for order 5020")
        self.assertEqual(subjects["cancel"], "Your request to cancel or change order 5020")
        r = self.ask("status", 5020, D(10, 12))
        copied = ui.full_email(r["subject"], r["email"])
        self.assertTrue(copied.startswith("Subject: Your order 5020: delivery update\n\nHi Miles,\n\n"))
        self.assertTrue(copied.endswith("Best,\nStark Industries Support"))
        self.assertEqual(ui.full_email("", r["email"]), r["email"])

    def test_repeat_inquiry_reports_the_existing_order(self):
        p = temp_book()
        first = save(self.ask("lost", 1650, D(10, 30), path=p), p)
        number = first["saved_orders"][0]["number"]
        again = self.ask("lost", 1650, D(11, 2), path=p)
        self.assertEqual(again["previews"], [])
        self.assertIn(f"We already sent a replacement for order 1650 under order number {number}.", again["email"])
        self.assertEqual(again["status"]["label"], f"Replacement already sent (order {number})")
        self.assertIn(f"order {number} already covers this", again["recommendation"]["title"])
        self.assertEqual(again["email_issues"], [])


class ShippingRules(unittest.TestCase):
    """Each kind of new order uses its own shipping method, and that method sets its dates.
    Expected dates were counted by hand on a calendar (application date Oct 8, 2026)."""

    CASES = [  # inquiry, order (its own method), contact date, fields, new method, ships, due, contact-us
        ("lost", 1024, "Express", D(10, 30), {}, "Regular", D(10, 30), D(11, 3), D(11, 6)),      # New York: X = 2 business days
        ("lost", 1650, "Regular", D(10, 30), {}, "Regular", D(10, 30), D(11, 5), D(11, 12)),     # Johannesburg: X = 4
        ("damaged", 2314, "Regular", D(10, 9), {"qty": 1}, "Express", D(10, 9), D(10, 11), D(10, 14)),   # Toronto: 2 calendar days
        ("wrong", 3520, "Regular", D(10, 12), {}, "Express", D(10, 12), D(10, 13), D(10, 15)),   # Montreal: 1 calendar day
        ("damaged", 3514, "Express", D(10, 8), {}, "Express", D(10, 8), D(10, 9), D(10, 11)),    # New York: 1 calendar day
        ("exchange_tracking", 1628, "Regular", D(10, 12),
         {"qty": 2, "new_item": "Waterproof Spray (Medium)", "tracking_number": "1Z9", "offered_on": D(10, 8)},
         "Regular", D(10, 12), D(10, 14), D(10, 17)),                                           # New Jersey: X = 2
    ]

    def test_method_and_dates_for_every_kind_of_new_order(self):
        for kind, order, original, on, fields, method, ships, due, contact in self.CASES:
            p = temp_book()
            before = sha(p)
            r = ui.prepare_reply(kind, str(order), on, APP_A, path=p, **fields)
            self.assertEqual(r["order"]["method"], original, f"{kind} {order}")
            row = r["previews"][0]
            self.assertEqual((row["method"], row["ship"], row["est"], row["contact"]), (method, ships, due, contact),
                             f"{kind} {order}")
            self.assertEqual(sha(p), before)                         # the preview wrote nothing
            done = save(r, p)
            saved = done["saved_orders"][0]
            self.assertEqual((saved["method"], saved["ship"], saved["est"], saved["contact"]), (method, ships, due, contact))
            self.assertEqual(len(new_order_rows(p)), 1)
            ws = openpyxl.load_workbook(p)["New Orders"]
            self.assertEqual(ws["K2"].value, method)                 # recorded in the workbook

    def test_original_orders_keep_the_method_they_were_placed_with(self):
        p = temp_book()
        express = ui.check_order("3514", D(10, 7), APP_A, path=p)     # Express to New York: 1 calendar day
        regular = ui.check_order("5020", D(10, 12), APP_A, path=p)    # Regular to New York: 2 business days
        self.assertEqual((express["order"]["method"], dict(express["status"]["dates"])["Expected delivery"]),
                         ("Express", "Wednesday, October 7, 2026"))
        self.assertEqual((regular["order"]["method"], dict(regular["status"]["dates"])["Expected delivery"]),
                         ("Regular", "Tuesday, October 13, 2026"))


class ContactUsDate(unittest.TestCase):
    """The contact-us date (2X + 1) is shown in the result, explained internally, and is the same everywhere."""

    def test_two_x_plus_one_counted_by_hand(self):
        p = temp_book()
        for order, on, ship, due, two_x, contact in (
                (1650, D(10, 8), "Wednesday, September 30, 2026", "Tuesday, October 6, 2026",
                 "Monday, October 12, 2026", "Tuesday, October 13, 2026"),      # Regular, X = 4: 8 business days, then +1
                (3520, D(10, 8), "Thursday, October 8, 2026", "Tuesday, October 13, 2026",
                 "Friday, October 16, 2026", "Saturday, October 17, 2026"),     # Regular, X = 3: the +1 lands on a Saturday
                (1614, D(9, 26), "Thursday, September 24, 2026", "Sunday, September 27, 2026",
                 "Wednesday, September 30, 2026", "Thursday, October 1, 2026"),  # Express, X = 3 calendar days
                (1545, D(9, 21), "Saturday, September 19, 2026", "Tuesday, September 22, 2026",
                 "Thursday, September 24, 2026", "Friday, September 25, 2026")):  # Regular, shipped on a Saturday
            r = ui.check_order(str(order), on, APP_A, path=p)
            dates = dict(r["status"]["dates"])
            self.assertEqual((dates["Shipped"], dates["Expected delivery"], dates["Contact-us date"]), (ship, due, contact))
            lines = "\n".join(r["explanation"][0]["lines"])
            self.assertIn(f"The shipping date, {ship}, is day zero. Counting begins the following day.", lines)
            self.assertIn(f"after the shipping date: {two_x}.", lines)
            self.assertIn(f"Plus 1 calendar day gives the contact-us date: {contact}.", lines)

    def test_explanation_shows_x_two_x_and_the_days_counted(self):
        r = ui.check_order("1650", D(10, 8), APP_A, path=temp_book())
        self.assertEqual(r["explanation"][0]["title"], "Contact-us date for order 1650")
        self.assertEqual(r["explanation"][0]["lines"], [
            "Shipping method and destination: Regular shipping to Johannesburg.",
            "X for Regular shipping to Johannesburg is 4 business days (Shipping estimates sheet). "
            "Business days skip Saturdays and Sundays.",
            "The shipping date, Wednesday, September 30, 2026, is day zero. Counting begins the following day.",
            "Expected delivery is X, 4 business days after the shipping date: Tuesday, October 6, 2026.",
            "2X is 8 business days after the shipping date: Monday, October 12, 2026.",
            "Days counted: day 1 Thu Oct 1, 2026, day 2 Fri Oct 2, 2026, day 3 Mon Oct 5, 2026, day 4 Tue Oct 6, 2026, "
            "day 5 Wed Oct 7, 2026, day 6 Thu Oct 8, 2026, day 7 Fri Oct 9, 2026, day 8 Mon Oct 12, 2026.",
            "Plus 1 calendar day gives the contact-us date: Tuesday, October 13, 2026."])
        express = ui.check_order("1614", D(9, 26), APP_A, path=temp_book())["explanation"][0]["lines"]
        self.assertIn("X for Express shipping to Albuquerque is 3 calendar days (Shipping estimates sheet). "
                      "Calendar days include Saturdays and Sundays.", express)
        self.assertIn("Days counted: day 1 Fri Sep 25, 2026, day 2 Sat Sep 26, 2026, day 3 Sun Sep 27, 2026, "
                      "day 4 Mon Sep 28, 2026, day 5 Tue Sep 29, 2026, day 6 Wed Sep 30, 2026.", express)

    def test_new_order_gets_its_own_explanation(self):
        r = ui.prepare_reply("lost", "1024", D(10, 30), APP_A, path=temp_book())
        self.assertEqual([e["title"] for e in r["explanation"]],
                         ["Contact-us date for order 1024", "Contact-us date for new order 5021 (preview)"])
        new = r["explanation"][1]["lines"]
        self.assertEqual(new[0], "A reshipment always goes by Regular shipping, whatever the original used. "
                                 "It ships the day it is created.")
        self.assertIn("X for Regular shipping to New York is 2 business days (Shipping estimates sheet). "
                      "Business days skip Saturdays and Sundays.", new)
        self.assertIn("Plus 1 calendar day gives the contact-us date: Friday, November 6, 2026.", new)
        self.assertIn("X for Express shipping to New York is 1 calendar day", "\n".join(r["explanation"][0]["lines"]))

    def test_same_dates_in_the_result_the_details_and_the_email(self):
        p = temp_book()
        for order in ORDERS:
            for days in (0, 2, 5):
                r = ui.check_order(str(order), APP_A + timedelta(days), APP_A, path=p)
                if not r["email"]:
                    continue
                shown, table = dict(r["status"]["dates"]), dict(r["key_dates"])
                f = r["internal"]["facts"]
                contact = date.fromisoformat(f["contact_us_date"])
                self.assertEqual(shown["Contact-us date"], ui.show_long(contact))
                self.assertEqual(table["Contact-us date"], ui.show(contact))
                self.assertEqual(shown["Expected delivery"], ui.show_long(date.fromisoformat(f["expected_delivery"])))
                if r["decision"] != E.RESHIP:
                    self.assertIn(E.fmt(contact), r["email"], f"{order} +{days}")       # the customer is told the date
                    self.assertIn(ui.show_long(contact), r["status"]["advice"])
                self.assertIn(ui.show_long(contact), "\n".join(r["explanation"][0]["lines"]))

    def test_new_order_dates_match_between_preview_saved_order_and_email(self):
        p = temp_book()
        preview = ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p)
        details = dict(preview["previews"][0]["details"])
        self.assertEqual((details["Ships"], details["Expected delivery"], details["Contact-us date"], details["Item"]),
                         ("Friday, October 30, 2026", "Thursday, November 5, 2026", "Thursday, November 12, 2026",
                          "1 × vibranium plate, red"))
        self.assertIn("5021", details["New order number"])
        done = save(preview, p)
        saved = dict(done["saved_orders"][0]["details"])
        self.assertEqual(saved["New order number"], "5021")
        self.assertEqual({k: saved[k] for k in ("Ships", "Expected delivery", "Contact-us date")},
                         {k: details[k] for k in ("Ships", "Expected delivery", "Contact-us date")})
        self.assertIn("Your replacement order number is 5021. It's scheduled to ship on Friday, October 30 and should "
                      "arrive by Thursday, November 5. If it hasn't arrived by Thursday, November 12, please contact us "
                      "again so we can help resolve the issue.", done["email"])


def sheet_values(path):
    """Every cell value in every sheet, formulas as text: {sheet: {cell: value}}."""
    wb = openpyxl.load_workbook(path)
    return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb}


def saved_rows(path):
    ws = openpyxl.load_workbook(path)["New Orders"]
    return [tuple(ws.cell(r, c).value for c in range(1, 13)) for r in new_order_rows(path)]


class ClearReplacement(unittest.TestCase):
    """Demonstration option: remove one saved reshipment or replacement so its original can be processed again."""

    def book_with_four_saved_orders(self):
        """5021 reship of 1650, 5022 damaged replacement for 3514, 5023 exchange for 1628, 5024 reship of 1024."""
        p = temp_book()
        for kind, order, on, fields in (
                ("lost", "1650", D(10, 30), {}), ("damaged", "3514", D(10, 8), {}),
                ("exchange_tracking", "1628", D(10, 12), {"qty": 2, "new_item": "Waterproof Spray (Medium)",
                                                          "tracking_number": "1Z9", "offered_on": D(10, 8)}),
                ("lost", "1024", D(10, 30), {})):
            self.assertTrue(save(ui.prepare_reply(kind, order, on, APP_A, path=p, **fields), p)["saved_orders"])
        self.assertEqual([r[:2] for r in saved_rows(p)], [(5021, 1650), (5022, 3514), (5023, 1628), (5024, 1024)])
        return p

    def test_clears_only_the_selected_replacement(self):
        p = self.book_with_four_saved_orders()
        before, rows = sheet_values(p), saved_rows(p)
        done = ui.clear_replacement(1650, 5021, path=p)
        self.assertTrue(done["ok"])
        self.assertEqual((done["cleared"], done["original"]), (5021, 1650))
        self.assertIn("Saved replacement order 5021 was cleared", done["message"])
        self.assertIn("No new order was created", done["message"])
        self.assertEqual(saved_rows(p), rows[1:])                    # the three unrelated saved orders, value for value
        after = sheet_values(p)
        for sheet in ("Orders", "Policies", "Shipping estimates", "Price List", "Orders 2"):
            self.assertEqual(after[sheet], before[sheet], sheet)     # originals, rolling formulas and every other sheet
        self.assertEqual((after["Orders"]["A6"], after["Orders"]["G6"]), (1650, "=G13-9"))   # order 1650 itself still rolls
        self.assertEqual({k: v for k, v in after["New Orders"].items() if k.endswith("1") and len(k) == 2},
                         {k: v for k, v in before["New Orders"].items() if k.endswith("1") and len(k) == 2})   # headers

    def test_refuses_everything_else_and_changes_nothing(self):
        p = self.book_with_four_saved_orders()
        before = sha(p)
        for original, number, expect in (
                (1628, 5023, "exchange order"),                      # an exchange is not covered
                (1024, 5021, "was not created from order 1024"),     # wrong original for that replacement
                (1650, 1650, "original order"),                      # an original order
                (1650, 3514, "original order"),
                (1650, 7777, "not a saved order")):                  # no such order
            r = ui.clear_replacement(original, number, path=p)
            self.assertFalse(r["ok"], (original, number))
            self.assertIn(expect, r["error"])
            self.assertEqual(sha(p), before)
        from stark_support import workbook as W
        for original, number in ((1628, 5023), (1024, 5021), (1650, 1650), (1650, 7777)):   # the writer checks again itself
            with self.assertRaises(ValueError):
                W.remove_new_order(p, number, original)
            self.assertEqual(sha(p), before)
        self.assertEqual(len(saved_rows(p)), 4)

    def test_only_reshipments_and_replacements_are_offered(self):
        p = self.book_with_four_saved_orders()
        self.assertEqual([(c["number"], c["original"]) for c in ui.check_order("1650", D(10, 30), APP_A, path=p)["clearable"]],
                         [(5021, 1650)])
        self.assertEqual([c["number"] for c in ui.prepare_reply("damaged", "3514", D(10, 8), APP_A, path=p)["clearable"]],
                         [5022])
        exchange = ui.check_order("1628", D(10, 13), APP_A, path=p)
        self.assertEqual([n["number"] for n in exchange["linked_orders"]], [5023])
        self.assertEqual(exchange["clearable"], [])                  # the exchange is listed but never clearable
        self.assertEqual(ui.check_order("5021", D(11, 2), APP_A, path=p)["clearable"], [])   # nor from the new order itself
        self.assertEqual(ui.check_order("5020", D(10, 12), APP_A, path=p)["clearable"], [])  # nothing saved for 5020

    def test_original_can_be_processed_again_and_nothing_is_saved_automatically(self):
        p = self.book_with_four_saved_orders()
        again = ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p)
        self.assertEqual((again["existing_order"], again["previews"]), (5021, []))
        self.assertEqual(again["status"]["label"], "Replacement already sent (order 5021)")
        self.assertTrue(ui.clear_replacement(1650, 5021, path=p)["ok"])
        cleared = sha(p)
        r = ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p)
        self.assertIsNone(r["existing_order"])
        self.assertEqual(r["status"]["label"], "Eligible for reshipment")
        self.assertEqual(r["decision"], E.RESHIP)
        row = r["previews"][0]
        self.assertEqual((row["number"], row["method"], row["ship"], row["est"], row["contact"]),
                         (5025, "Regular", D(10, 30), D(11, 5), D(11, 12)))        # a new number; same rules, same dates
        self.assertTrue(r["email_withheld"])
        self.assertEqual(sha(p), cleared)                             # checking again saved nothing
        self.assertEqual(len(saved_rows(p)), 3)
        self.assertEqual(ui.prepare_reply("lost", "1650", D(10, 8), APP_A, path=p)["decision"], E.LATE)   # rules unchanged
        done = save(r, p)
        self.assertEqual(done["saved_orders"][0]["number"], 5025)
        self.assertEqual([x[:2] for x in saved_rows(p)], [(5022, 3514), (5023, 1628), (5024, 1024), (5025, 1650)])

    def test_cleared_number_is_never_issued_again(self):
        """Clear the highest-numbered order, create another: it gets a new, unused number, not the cleared one."""
        p = temp_book()
        self.assertEqual(save(ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p), p)["saved_orders"][0]["number"], 5021)
        self.assertTrue(ui.clear_replacement(1650, 5021, path=p)["ok"])
        self.assertEqual(saved_rows(p), [])
        cleared = sha(p)
        preview = ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p)
        self.assertEqual(preview["previews"][0]["number"], 5022)      # the preview already shows the new number
        self.assertEqual(sha(p), cleared)
        done = save(preview, p)
        self.assertEqual(done["saved_orders"][0]["number"], 5022)
        self.assertIn("Your replacement order number is 5022.", done["email"])
        self.assertEqual([r[:2] for r in saved_rows(p)], [(5022, 1650)])
        # again and again: every round gets a number that has never been used
        issued = [5021, 5022]
        for _ in range(3):
            self.assertTrue(ui.clear_replacement(1650, issued[-1], path=p)["ok"])
            issued.append(save(ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p), p)["saved_orders"][0]["number"])
        self.assertEqual(issued, [5021, 5022, 5023, 5024, 5025])
        self.assertEqual(len(set(issued)), len(issued))

    def test_other_kinds_of_order_also_skip_cleared_numbers(self):
        p = temp_book()
        save(ui.prepare_reply("damaged", "3514", D(10, 8), APP_A, path=p), p)        # 5021
        save(ui.prepare_reply("lost", "1024", D(10, 30), APP_A, path=p), p)          # 5022
        self.assertTrue(ui.clear_replacement(1024, 5022, path=p)["ok"])             # clear the highest
        exchange = save(ui.prepare_reply("exchange_tracking", "1628", D(10, 12), APP_A, path=p, qty=2,
                                         new_item="Waterproof Spray (Medium)", tracking_number="1Z9",
                                         offered_on=D(10, 8)), p)
        self.assertEqual(exchange["saved_orders"][0]["number"], 5023)                # not 5022
        self.assertTrue(ui.clear_replacement(3514, 5021, path=p)["ok"])             # clear a lower one
        self.assertEqual(save(ui.prepare_reply("damaged", "3514", D(10, 8), APP_A, path=p), p)["saved_orders"][0]["number"],
                         5024)                                                       # not 5021, not 5022
        self.assertEqual([r[:2] for r in saved_rows(p)], [(5023, 1628), (5024, 3514)])

    def test_issued_number_is_recorded_in_the_workbook_and_only_goes_up(self):
        from stark_support import workbook as W
        p = temp_book()
        cell = lambda: openpyxl.load_workbook(p)["New Orders"][W.LAST_NUMBER_CELL].value  # noqa: E731
        self.assertIsNone(cell())
        self.assertEqual(ui.load(p, app_date=APP_A).last_number, 0)
        self.assertEqual(ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p)["previews"][0]["number"], 5021)
        self.assertIsNone(cell())                                     # a preview records nothing
        save(ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p), p)
        save(ui.prepare_reply("lost", "1024", D(10, 30), APP_A, path=p), p)
        self.assertEqual(cell(), 5022)
        self.assertEqual(openpyxl.load_workbook(p)["New Orders"][W.LAST_NUMBER_LABEL_CELL].value, W.LAST_NUMBER_LABEL)
        self.assertTrue(ui.clear_replacement(1650, 5021, path=p)["ok"])             # clearing a lower number
        self.assertEqual(cell(), 5022)
        self.assertTrue(ui.clear_replacement(1024, 5022, path=p)["ok"])             # clearing the highest
        self.assertEqual(cell(), 5022)                                # still remembered with no orders left
        self.assertEqual(ui.load(p, app_date=APP_A).last_number, 5022)
        self.assertEqual(saved_rows(p), [])
        self.assertEqual([c.value for c in openpyxl.load_workbook(p)["New Orders"][1][:12]], W.NEW_ORDER_HEADERS)

    def test_workbook_saved_before_numbers_were_recorded(self):
        """Like the working workbook today: saved orders 5021 and 5022 exist, but no number has been recorded yet."""
        from stark_support import workbook as W
        p = temp_book()
        save(ui.prepare_reply("lost", "1650", D(10, 30), APP_A, path=p), p)
        save(ui.prepare_reply("lost", "1024", D(10, 30), APP_A, path=p), p)
        wb = openpyxl.load_workbook(p)
        wb["New Orders"][W.LAST_NUMBER_LABEL_CELL] = wb["New Orders"][W.LAST_NUMBER_CELL] = None
        wb.save(p)
        self.assertEqual(ui.load(p, app_date=APP_A).last_number, 0)
        self.assertTrue(ui.clear_replacement(1024, 5022, path=p)["ok"])             # the highest-numbered order
        self.assertEqual(ui.load(p, app_date=APP_A).last_number, 5022)              # recorded at the moment it is cleared
        self.assertEqual(save(ui.prepare_reply("lost", "1024", D(10, 30), APP_A, path=p), p)["saved_orders"][0]["number"],
                         5023)
        self.assertEqual([r[:2] for r in saved_rows(p)], [(5021, 1650), (5023, 1024)])

    def test_on_an_exact_copy_of_the_working_workbook(self):
        """The real file as it is now, copied: clear its highest-numbered reshipment or replacement, create it again."""
        from stark_support import workbook as W
        p = os.path.join(TMP, f"live{len(os.listdir(TMP))}.xlsx")
        shutil.copy(BOOK, p)
        rows = saved_rows(p)
        highest = max(rows, key=lambda r: r[0], default=None)
        if highest is None or highest[2] not in W.CLEARABLE_TYPES or highest[2] != W.NEW_ORDER_TYPES["lost"]:
            self.skipTest("the working workbook's highest-numbered saved order is not a reshipment")
        number, original = highest[0], highest[1]
        used = {r[0] for r in rows} | {o.number for o in ui.load(p, app_date=APP_A).orders}
        self.assertTrue(ui.clear_replacement(original, number, path=p)["ok"])
        self.assertEqual(saved_rows(p), [r for r in rows if r[0] != number])        # every other saved order as it was
        again = save(ui.prepare_reply("lost", str(original), D(12, 30), APP_A, path=p), p)
        new = again["saved_orders"][0]["number"]
        self.assertEqual(new, number + 1)
        self.assertNotIn(new, used)

    def test_other_workflows_are_unaffected_after_a_clear(self):
        p = self.book_with_four_saved_orders()
        self.assertTrue(ui.clear_replacement(3514, 5022, path=p)["ok"])             # a damaged-item replacement
        damaged = ui.prepare_reply("damaged", "3514", D(10, 8), APP_A, path=p)
        self.assertEqual((damaged["decision"], damaged["previews"][0]["method"], damaged["previews"][0]["est"]),
                         ("DM_CREATED", "Express", D(10, 9)))
        exchange = ui.prepare_reply("exchange_tracking", "1628", D(10, 13), APP_A, path=p, qty=2,
                                    new_item="Waterproof Spray (Medium)", tracking_number="1Z9", offered_on=D(10, 8))
        self.assertEqual(exchange["existing_order"], 5023)            # the exchange is still there and still guarded
        refund = ui.prepare_reply("refund", "1814", D(10, 8), APP_A, path=p, qty=1)
        self.assertEqual((refund["decision"], refund["internal"]["facts"]["refund_amount"]), ("REFUND_OK", 15))
        self.assertEqual(ui.prepare_reply("lost", "1024", D(11, 2), APP_A, path=p)["existing_order"], 5024)

    def test_locked_workbook_gives_the_close_excel_message(self):
        p = self.book_with_four_saved_orders()
        before = sha(p)
        with mock.patch("stark_support.workbook.remove_new_order", side_effect=PermissionError(13, "Permission denied")):
            with self.assertLogs("stark_support.ui", level="ERROR"):
                r = ui.clear_replacement(1650, 5021, path=p)
        self.assertEqual(r["error"], f"Close {os.path.basename(p)} in Excel and click Confirm clear again.")
        self.assertEqual(sha(p), before)


class Saving(unittest.TestCase):
    def test_double_click_saves_exactly_one_row(self):
        p = temp_book()
        state = {}
        check = ui.check_order("1650", D(10, 30), APP_A, path=p)
        token = ui.remember_pending(state, check)
        first = ui.save_once(state, token, path=p)
        second = ui.save_once(state, token, path=p)                 # double click / rerun with the same token
        self.assertTrue(first["ok"])
        number = first["saved_orders"][0]["number"]
        self.assertEqual(number, 5021)
        self.assertIn(str(number), first["email"])
        self.assertTrue(second["repeat"])
        self.assertEqual(second["saved_orders"][0]["number"], number)
        self.assertEqual(len(new_order_rows(p)), 1)
        self.assertEqual(state["pending"], {})

    def test_saving_the_same_request_again_returns_the_existing_order(self):
        p = temp_book()
        check = ui.check_order("1650", D(10, 30), APP_A, path=p)
        first = ui.save_action(check["request"]["params"], APP_A, path=p)
        again = ui.save_action(check["request"]["params"], APP_A, path=p)   # a fresh token for the same request
        number = first["saved_orders"][0]["number"]
        self.assertEqual(again["saved_orders"], [])
        self.assertEqual(again["existing_order"], number)
        self.assertTrue(again["decision"].startswith("ALREADY_RESHIPPED"))
        self.assertEqual(len(new_order_rows(p)), 1)
        later = ui.check_order("1650", D(10, 30), APP_A, path=p)
        self.assertEqual(later["existing_order"], number)
        self.assertEqual(later["previews"], [])
        self.assertEqual([(n["number"], n["created"], n["method"]) for n in later["linked_orders"]],
                         [(number, D(10, 30), "Regular")])

    def test_saved_number_comes_from_the_save_not_the_preview(self):
        p = temp_book()
        state = {}
        a = ui.check_order("1650", D(10, 30), APP_A, path=p)
        b = ui.run_action("damaged", 3514, D(10, 8), APP_A, path=p)
        self.assertEqual(a["previews"][0]["number"], b["previews"][0]["number"])   # both previews say 5021
        ta, tb = ui.remember_pending(state, a), ui.remember_pending(state, b)
        self.assertEqual(ui.save_once(state, tb, path=p)["saved_orders"][0]["number"], 5021)
        self.assertEqual(ui.save_once(state, ta, path=p)["saved_orders"][0]["number"], 5022)
        self.assertEqual(len(new_order_rows(p)), 2)

    def test_permission_error_becomes_the_close_excel_message(self):
        p = temp_book()
        before = sha(p)
        state = {}
        token = ui.remember_pending(state, ui.check_order("1650", D(10, 30), APP_A, path=p))
        with mock.patch("stark_support.workbook.append_new_order", side_effect=PermissionError(13, "Permission denied")):
            with self.assertLogs("stark_support.ui", level="ERROR"):
                r = ui.save_once(state, token, path=p)
        self.assertFalse(r["ok"])
        self.assertEqual(r["error"], f"Close {os.path.basename(p)} in Excel and click Create order again.")
        self.assertEqual(sha(p), before)
        self.assertIn(token, state["pending"])                     # nothing is lost: the same click works afterwards
        self.assertTrue(ui.save_once(state, token, path=p)["ok"])
        self.assertEqual(len(new_order_rows(p)), 1)

    def test_default_message_names_the_real_workbook(self):
        with mock.patch("stark_support.ui_adapter.load", side_effect=PermissionError(13, "Permission denied")):
            with self.assertLogs("stark_support.ui", level="ERROR"):
                r = ui.save_action({"kind": "lost", "order_number": 1650, "on": D(10, 30)}, APP_A)
        self.assertEqual(r["error"], f"Close {os.path.basename(ui.WORKBOOK_PATH)} in Excel and click Create order again.")


class Page(unittest.TestCase):
    """Headless runs of app.py against a temporary copy of the workbook."""

    def open(self, path):
        from streamlit.testing.v1 import AppTest
        patcher = mock.patch.object(ui, "WORKBOOK_PATH", path)
        patcher.start()
        self.addCleanup(patcher.stop)
        at = AppTest.from_file(os.path.join(ROOT, "app.py"), default_timeout=60).run()
        self.assertFalse(at.exception)
        return at

    def fill(self, at, order, on, kind="status", app=APP_A):
        at.text_input(key="ident").set_value(order)
        at.selectbox(key="inquiry").set_value(kind)
        at.date_input(key="contact").set_value(on)
        at.date_input(key="app").set_value(app)
        at.run()
        return at

    def check(self, at):
        at.button(key="check").click().run()
        self.assertFalse(at.exception)
        return at

    @staticmethod
    def text(at):
        return "\n".join(m.value for m in at.markdown)

    def test_starts_with_the_inquiry_and_one_main_button(self):
        at = self.open(temp_book())
        self.assertEqual(at.title[0].value, "STARK INDUSTRIES Order Support")
        self.assertEqual(at.button(key="check").label, "Check Order and Prepare Reply")
        self.assertEqual([w.label for w in (at.text_input(key="ident"), at.selectbox(key="inquiry"),
                                            at.date_input(key="contact"), at.date_input(key="app"))],
                         ["Order number", "What is the customer asking?", "Date the customer contacted us",
                          "Application date (normally today)"])
        self.assertIsNone(at.date_input(key="contact").value)       # never guessed
        self.assertEqual(at.date_input(key="app").value, ui.local_today())
        self.assertEqual(len(at.sidebar), 0)

    def test_validation_messages(self):
        p = temp_book()
        before = sha(p)
        at = self.open(p)
        self.check(at)
        self.assertIn("date the customer contacted you", at.error[0].value)
        self.check(self.fill(at, "", D(10, 8)))
        self.assertIn("Enter an order number", at.error[0].value)
        self.check(self.fill(at, "sixteen fifty", D(10, 8)))
        self.assertIn("digits only", at.error[0].value)
        self.check(self.fill(at, "9999", D(10, 8)))
        self.assertIn("No order with the number 9999", at.error[0].value)
        self.check(self.fill(at, "1650", D(10, 8), app=APP_B))
        self.assertIn("earlier than order 1650's order date", at.warning[0].value)
        self.assertEqual(len(at.text_area), 0)                      # no email on any of these
        self.assertEqual(sha(p), before)

    def test_result_order_status_action_email_details_internal(self):
        p = temp_book()
        before = sha(p)
        at = self.check(self.fill(self.open(p), "1650", D(10, 8)))
        self.assertEqual([h.value for h in at.subheader],
                         ["Customer inquiry", "1. Order status", "2. Recommended action", "3. Customer email"])
        self.assertEqual([e.label for e in at.expander],
                         ["4. Order details", "5. Internal explanation (not for the customer)"])
        page = self.text(at)
        self.assertIn("Running late, waiting for reshipment eligibility", page)
        self.assertIn("**Preview** · Order 1650, M'Baku · Where is my order? · contacted Thu Oct 8, 2026 · "
                      "application date Thu Oct 8, 2026", page)
        self.assertIn("Do not reship yet, send the reply", at.info[0].value)
        self.assertEqual(at.text_input(key="email_subject").value, "Your order 1650: delivery update")
        body = at.text_area(key="email_body").value
        self.assertIn("I'm sorry, your order 1650 is running late.", body)
        self.assertEqual(at.button(key="copy_email").label, "Copy Email")
        self.assertNotIn("decision", body.lower())                  # internal reasoning stays out of the email
        self.assertEqual(len(at.warning), 0)
        self.assertEqual(sha(p), before)

    def test_inquiry_changes_the_email_without_restart(self):
        p = temp_book()
        before = sha(p)
        at = self.check(self.fill(self.open(p), "1814", D(10, 8)))
        self.assertIn("is running late", at.text_area(key="email_body").value)
        at.selectbox(key="inquiry").set_value("refund").run()
        self.assertIn("You changed the inquiry", at.warning[0].value)
        at.number_input(key="qty").set_value(1).run()
        self.check(at)
        self.assertIn("Yes, you can return the Taser (High Voltage) for a refund.", at.text_area(key="email_body").value)
        self.assertEqual(at.text_input(key="email_subject").value, "Your refund request for order 1814")
        self.assertIn("Refund allowed: $15 voucher after the return", at.success[0].value.replace("\\$", "$"))
        self.check(self.fill(at, "1650", D(10, 8), kind="cancel"))
        self.assertIn("Change or cancellation unavailable", at.info[0].value)
        self.assertIn("aren't able to cancel or change", at.text_area(key="email_body").value)
        self.assertEqual(sha(p), before)

    def test_edited_email_is_checked_before_copying(self):
        at = self.check(self.fill(self.open(temp_book()), "1650", D(10, 8)))
        self.assertIn("Checked against this order", "\n".join(c.value for c in at.caption))
        at.text_area(key="email_body").set_value("Hi M'Baku,\n\nYour order is lost. It arrives October 20.\n\nBest,\nSam").run()
        warning = at.warning[0].value
        self.assertIn("internal terms", warning)
        self.assertIn("October 20", warning)

    def test_email_search_lets_me_pick_the_order(self):
        at = self.check(self.fill(self.open(temp_book()), "multi-order-b@example.com", D(10, 8), kind="refund"))
        self.assertEqual([b.key for b in at.button if str(b.key).startswith("match_")],
                         ["match_1655", "match_1855", "match_2314"])
        at.button(key="match_1655").click().run()
        self.assertFalse(at.exception)
        self.assertEqual(at.text_input(key="ident").value, "1655")
        self.assertIn("Hoverboards can't be returned for a refund", at.text_area(key="email_body").value)

    def test_preview_never_saves_and_create_order_saves_once(self):
        p = temp_book()
        before = sha(p)
        at = self.check(self.fill(self.open(p), "1650", D(10, 30), kind="lost"))
        self.assertIn("Create the replacement order (Regular shipping)", at.warning[0].value)
        self.assertIn("Replacement order: preview only, not saved.", at.warning[1].value)
        page = self.text(at)
        for line in ("**New order number:** 5021", "**Item:** 1 × vibranium plate, red", "**Ships:** Friday, October 30, 2026",
                     "**Expected delivery:** Thursday, November 5, 2026", "**Contact-us date:** Thursday, November 12, 2026",
                     "**Shipping method (internal, not in the email):** Regular"):
            self.assertIn(line, page)
        self.assertIn("The email appears after you create the order", at.info[0].value)
        self.assertEqual(len(at.text_area), 0)
        self.assertEqual([b.key for b in at.button if b.key in ("create", "save", "cancel")], ["create"])
        at.button(key="create").click().run()                       # confirm step: still nothing written
        self.assertEqual(sha(p), before)
        self.assertIn("Save this replacement order to the workbook?", self.text(at))
        self.assertEqual([(b.key, b.label) for b in at.button if b.key in ("create", "save", "cancel")],
                         [("save", "Confirm and save"), ("cancel", "Cancel")])
        token = at.session_state["token"]
        at.button(key="save").click().run()
        self.assertFalse(at.exception)
        self.assertEqual(len(new_order_rows(p)), 1)
        self.assertIn("Done: order 5021 created", at.success[0].value)
        self.assertIn("Replacement order saved.", at.success[1].value)
        body = at.text_area(key="email_body").value
        self.assertIn("Your replacement order number is 5021. It's scheduled to ship on Friday, October 30", body)
        self.assertIn("**Saved** · Order 1650", self.text(at))
        self.assertNotIn("save", [b.key for b in at.button])
        state = {k: at.session_state[k] for k in ("pending", "saved")}
        self.assertTrue(ui.save_once(state, token, path=p)["repeat"])      # a replayed click
        self.check(at)                                              # asking again points to the existing order
        self.assertIn("order 5021 already covers this", at.info[0].value)
        self.assertEqual(len(new_order_rows(p)), 1)

    def test_cancel_saves_nothing_and_the_order_can_still_be_created(self):
        p = temp_book()
        before = sha(p)
        at = self.check(self.fill(self.open(p), "1024", D(10, 30), kind="lost"))
        at.button(key="create").click().run()
        at.button(key="cancel").click().run()
        self.assertFalse(at.exception)
        self.assertEqual(sha(p), before)                             # the workbook file is byte-identical
        self.assertEqual(new_order_rows(p), [])
        self.assertIn("Cancelled. Nothing was saved.", "\n".join(i.value for i in at.info))
        self.assertEqual([b.key for b in at.button if b.key in ("create", "save", "cancel")], ["create"])
        self.assertEqual(len(at.text_area), 0)                       # still no email for an order that doesn't exist
        at.button(key="create").click().run()                        # changing my mind still works, once
        at.button(key="save").click().run()
        self.assertEqual(len(new_order_rows(p)), 1)
        self.check(at)                                               # repeating the inquiry creates nothing more
        self.assertNotIn("create", [b.key for b in at.button])
        self.assertIn("order 5021 already covers this", at.info[0].value)
        self.assertEqual(len(new_order_rows(p)), 1)

    def test_clear_saved_replacement_needs_confirmation_and_restores_the_workflow(self):
        p = temp_book()
        at = self.check(self.fill(self.open(p), "1024", D(10, 30), kind="lost"))
        self.assertNotIn("clear_5021", [b.key for b in at.button])   # nothing saved yet, so nothing to clear
        at.button(key="create").click().run()
        at.button(key="save").click().run()
        self.assertEqual([r[:2] for r in saved_rows(p)], [(5021, 1024)])
        saved = sha(p)
        self.assertIn("Demonstration and testing: clear a saved replacement", [e.label for e in at.expander])
        self.assertEqual(at.button(key="clear_5021").label, "Clear saved replacement for retesting")
        self.assertNotIn("confirm_clear", [b.key for b in at.button])

        at.button(key="clear_5021").click().run()                    # asks first; nothing removed yet
        self.assertEqual(sha(p), saved)
        box = "\n".join(w.value for w in at.warning)
        for text in ("Clear saved replacement order 5021?", "Original order: 1024", "Saved replacement order: 5021",
                     "so that order 1024 can be processed again"):
            self.assertIn(text, box)
        self.assertEqual([(b.key, b.label) for b in at.button if "clear" in str(b.key)],
                         [("confirm_clear", "Confirm clear"), ("cancel_clear", "Cancel")])

        at.button(key="cancel_clear").click().run()                  # cancel: nothing changes
        self.assertEqual(sha(p), saved)
        self.assertIn("Cancelled. Nothing was changed.", "\n".join(i.value for i in at.info))
        self.assertEqual([b.key for b in at.button if "clear" in str(b.key)], ["clear_5021"])

        at.button(key="clear_5021").click().run()
        at.button(key="confirm_clear").click().run()                 # confirm: that one row goes
        self.assertFalse(at.exception)
        self.assertEqual(saved_rows(p), [])
        cleared = sha(p)
        self.assertIn("Saved replacement order 5021 was cleared from the workbook. Order 1024 can be processed again.",
                      "\n".join(x.value for x in at.success))
        page = self.text(at)
        self.assertIn("Eligible for reshipment", page)               # no longer "Replacement already sent"
        self.assertNotIn("Replacement already sent", page)
        self.assertIn("Replacement order: preview only, not saved.", "\n".join(w.value for w in at.warning))
        self.assertEqual(len(at.text_area), 0)                       # and no email until a new order is really saved
        self.assertEqual([b.key for b in at.button if b.key in ("create", "save", "cancel")], ["create"])
        self.assertNotIn("Demonstration and testing: clear a saved replacement", [e.label for e in at.expander])
        self.assertEqual(sha(p), cleared)                            # clearing did not create anything

        self.assertIn("**New order number:** 5022", page)             # a new number, not the cleared 5021
        at.button(key="create").click().run()                        # the normal workflow saves a new one
        self.assertEqual(sha(p), cleared)
        at.button(key="save").click().run()
        self.assertEqual([r[:2] for r in saved_rows(p)], [(5022, 1024)])
        self.assertIn("Your replacement order number is 5022.", at.text_area(key="email_body").value)

    def test_status_shows_the_three_dates_and_what_the_customer_should_do(self):
        at = self.check(self.fill(self.open(temp_book()), "1650", D(10, 8)))
        page = self.text(at)
        self.assertIn("**Shipped:** Wednesday, September 30, 2026  \n**Expected delivery:** Tuesday, October 6, 2026  \n"
                      "**Contact-us date:** Tuesday, October 13, 2026", page)
        self.assertIn("**If the package has not arrived by Tuesday, October 13, 2026, the customer should contact us again "
                      "so we can help resolve the issue.**", page)
        self.assertIn("Tuesday, October 13", at.text_area(key="email_body").value)
        self.assertIn("Plus 1 calendar day gives the contact-us date: Tuesday, October 13, 2026.", page)
        self.assertNotIn("2X", at.text_area(key="email_body").value)

    def test_check_another_order_clears_the_page(self):
        at = self.check(self.fill(self.open(temp_book()), "1650", D(10, 8)))
        at.button(key="another").click().run()
        self.assertFalse(at.exception)
        self.assertEqual(at.text_input(key="ident").value, "")
        self.assertIsNone(at.date_input(key="contact").value)
        self.assertEqual(len(at.text_area), 0)
        self.assertIn("Fill in the customer inquiry above", at.info[0].value)

    def test_check_another_order_resets_every_field_and_keeps_saved_orders(self):
        p = temp_book()
        at = self.fill(self.open(p), "1628", D(10, 12), kind="exchange_tracking", app=APP_A)
        at.text_input(key="item").set_value("quiver")
        at.number_input(key="qty").set_value(2)
        at.selectbox(key="new_item").set_value("Waterproof Spray (Medium)")
        at.text_input(key="tracking").set_value("1Z999")
        at.date_input(key="offered").set_value(D(10, 8))
        self.check(at.run())
        at.button(key="create").click().run()
        at.button(key="save").click().run()
        self.assertIn("Done: order 5021 created", at.success[0].value)
        saved = sha(p)
        at.button(key="another").click().run()
        self.assertFalse(at.exception)
        self.assertEqual((at.text_input(key="ident").value, at.selectbox(key="inquiry").value,
                          at.date_input(key="contact").value, at.date_input(key="app").value),
                         ("", "status", None, ui.local_today()))
        for key in ("result", "checked", "email_subject", "email_body"):
            self.assertNotIn(key, at.session_state)
        self.assertEqual(len(at.success) + len(at.warning) + len(at.error) + len(at.text_area), 0)
        at.selectbox(key="inquiry").set_value("exchange_tracking").run()        # the extra fields come back empty
        self.assertEqual((at.text_input(key="item").value, at.number_input(key="qty").value,
                          at.selectbox(key="new_item").value, at.text_input(key="tracking").value,
                          at.date_input(key="offered").value), ("", None, None, "", None))
        self.assertEqual(sha(p), saved)                              # the saved order is untouched
        self.assertEqual(len(new_order_rows(p)), 1)


class ZZDemoFilesUntouched(unittest.TestCase):
    """Named to run last: the committed demo workbook is unchanged."""

    def test_demo_workbook_unchanged_by_these_tests(self):
        self.assertEqual(sha(BOOK), START["book"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
